from openpyxl import load_workbook
from inventory.models import productsModel
from django.db import transaction
import logging
from datetime import datetime

class ExcelImportService:
    def __init__(self):
        self.logger = logging.getLogger(__name__)

    def inbound_excel_file_process(self, file_path):
        """
        Process Excel file and extract inbound data
        Returns JSON format compatible with inv_inboundProcess view

        Parameters:
        - file_path: Path to Excel file

        Excel Format Requirements:
        - sku: Unique product identifier (required)
        - product_name: Product name (required)
        - category: Product category (optional, defaults to 'Other')
        - unit_price: Purchase price per unit (required)
        - quantity: Number of units (required)
        - expiry_date: Expiration date (optional, YYYY-MM-DD format)
        """
        try:
            workbook = load_workbook(filename=file_path)
            sheet = workbook.active

            # Validate headers
            headers = [str(cell.value).strip().lower() if cell.value else '' for cell in sheet[1]]
            required_headers = ['sku', 'product_name', 'unit_price', 'quantity']
            self._validate_headers(headers, required_headers)

            # Prepare data for inbound registration
            items = []
            created_products = []

            # Get column indices
            col_indices = {
                'sku': headers.index('sku'),
                'product_name': headers.index('product_name'),
                'category': headers.index('category') if 'category' in headers else None,
                'unit_price': headers.index('unit_price'),
                'quantity': headers.index('quantity'),
                'expiry_date': headers.index('expiry_date') if 'expiry_date' in headers else None
            }

            for row in sheet.iter_rows(min_row=2, values_only=True):
                try:
                    # Extract data from row using column indices
                    sku = str(row[col_indices['sku']]).strip()
                    product_name = str(row[col_indices['product_name']]).strip()
                    category = str(row[col_indices['category']]).strip() if col_indices['category'] is not None and row[col_indices['category']] else 'Other'
                    unit_price = float(row[col_indices['unit_price']]) if row[col_indices['unit_price']] else 0.0
                    quantity = int(row[col_indices['quantity']]) if row[col_indices['quantity']] else 0

                    # Handle expiry_date if present
                    expiry_date = None
                    if col_indices['expiry_date'] is not None and row[col_indices['expiry_date']]:
                        try:
                            if isinstance(row[col_indices['expiry_date']], datetime):
                                expiry_date = row[col_indices['expiry_date']]
                            else:
                                expiry_date = datetime.strptime(str(row[col_indices['expiry_date']]), '%Y-%m-%d')
                        except ValueError:
                            # Try alternative date formats
                            try:
                                expiry_date = datetime.strptime(str(row[col_indices['expiry_date']]), '%m/%d/%Y')
                            except ValueError:
                                expiry_date = None

                    # Check if product exists by SKU only, create if not exists
                    product, created = self._get_or_create_product(
                        sku=sku,
                        name=product_name,
                        price=unit_price,
                        category=category
                    )

                    if created:
                        created_products.append(sku)
                        self.logger.info(f"Created product: {product_name} (SKU: {sku})")

                    # Add to items list in format expected by inv_inboundProcess
                    items.append({
                        'product': product.id,
                        'quantity': quantity,
                        'unit_price': unit_price,
                        'expiry_date': expiry_date.strftime('%Y-%m-%d') if expiry_date else None
                    })

                except Exception as row_error:
                    self.logger.error(f"Error processing row: {str(row_error)}")
                    continue

            if not items:
                raise ValueError("No valid items found in Excel file")

            # Calculate total price
            total_price = sum(item['unit_price'] * item['quantity'] for item in items)

            # Return data in same format as normal inbound registration
            # Use parameters from frontend form
            return {
                'success': True,
                'data': {
                    # 'supplier': supplier,
                    # 'ref_num': reference_number,
                    # 'total_price': total_price,
                    # 'inbound_date': inbound_date,
                    'items': items
                },
                'created_products': created_products,
                'message': f"Processed {len(items)} items, created {len(created_products)} new products"
            }

        except Exception as e:
            self.logger.error(f"Excel processing error: {str(e)}")
            return {
                'success': False,
                'error': str(e),
                'message': 'Failed to process Excel file'
            }

    def _validate_headers(self, headers, required_headers):
        """Validate Excel has required columns"""
        missing = [h for h in required_headers if h not in headers]
        if missing:
            raise ValueError(f"Missing required columns: {', '.join(missing)}")

    def outbound_excel_file_process(self, file_path):
        """
        Process Excel file and extract outbound data
        Returns JSON format compatible with inv_outboundProcess view

        Parameters:
        - file_path: Path to Excel file

        Excel Format Requirements:
        - sku: Unique product identifier (required)
        - quantity: Number of units to dispatch (required)
        """
        try:
            workbook = load_workbook(filename=file_path)
            sheet = workbook.active

            # Validate headers for outbound
            headers = [str(cell.value).strip().lower() if cell.value else '' for cell in sheet[1]]
            required_headers = ['sku', 'quantity']
            self._validate_headers(headers, required_headers)

            # Prepare data for outbound registration
            items = []

            # Get column indices
            col_indices = {
                'sku': headers.index('sku'),
                'quantity': headers.index('quantity')
            }

            for row in sheet.iter_rows(min_row=2, values_only=True):
                try:
                    # Extract data from row using column indices
                    sku = str(row[col_indices['sku']]).strip()
                    quantity = int(row[col_indices['quantity']]) if row[col_indices['quantity']] else 0

                    # Check if product exists by SKU
                    try:
                        product = productsModel.objects.get(sku=sku)
                        product_id = product.id
                    except productsModel.DoesNotExist:
                        raise ValueError(f"Product with SKU {sku} does not exist")

                    # Add to items list in format expected by inv_outboundProcess
                    items.append({
                        'product': product_id,
                        'quantity': quantity,
                    })

                except Exception as row_error:
                    self.logger.error(f"Error processing row: {str(row_error)}")
                    continue

            if not items:
                raise ValueError("No valid items found in Excel file")

            # Return data in same format as normal outbound registration
            return {
                'success': True,
                'data': {
                    'items': items
                },
                'message': f"Processed {len(items)} items for outbound"
            }

        except Exception as e:
            self.logger.error(f"Excel processing error: {str(e)}")
            return {
                'success': False,
                'error': str(e),
                'message': 'Failed to process Excel file'
            }

    def _validate_headers(self, headers, required_headers):
        """Validate Excel has required columns"""
        missing = [h for h in required_headers if h not in headers]
        if missing:
            raise ValueError(f"Missing required columns: {', '.join(missing)}")

    def product_excel_file_process(self, file_path):
        """
        Process Excel file and extract product data for bulk registration
        Returns JSON format with product data and processing results

        Parameters:
        - file_path: Path to Excel file

        Excel Format Requirements:
        - sku: Unique product identifier (required)
        - product_name: Product name (required)
        - category: Product category (required)
        - price: Product price (required)
        - threshold: Low stock threshold (required)
        - description: Product description (required)
        """
        try:
            workbook = load_workbook(filename=file_path)
            sheet = workbook.active

            # Validate headers
            headers = [str(cell.value).strip().lower() if cell.value else '' for cell in sheet[1]]
            required_headers = ['sku', 'product_name', 'category', 'price', 'threshold', 'description']
            self._validate_headers(headers, required_headers)

            # Prepare data for product creation
            products = []
            errors = []

            # Get column indices
            col_indices = {
                'sku': headers.index('sku'),
                'product_name': headers.index('product_name'),
                'category': headers.index('category'),
                'price': headers.index('price'),
                'threshold': headers.index('threshold'),
                'description': headers.index('description')
            }

            for row_idx, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
                try:
                    # Extract data from row using column indices
                    sku = str(row[col_indices['sku']]).strip()
                    product_name = str(row[col_indices['product_name']]).strip()
                    category = str(row[col_indices['category']]).strip()
                    price = float(row[col_indices['price']]) if row[col_indices['price']] else 0.0
                    threshold = int(row[col_indices['threshold']]) if row[col_indices['threshold']] else 0
                    description = str(row[col_indices['description']]).strip() if row[col_indices['description']] else 'No Description'

                    # Check if product with this SKU already exists
                    if productsModel.objects.filter(sku=sku).exists():
                        errors.append({
                            'row': row_idx,
                            'sku': sku,
                            'product_name': product_name,
                            'error': 'Duplicate SKU - product already exists'
                        })
                        continue

                    # Add to products list
                    products.append({
                        'sku': sku,
                        'name': product_name,
                        'category': category,
                        'price': int(price),  # Convert to int as per model
                        'description': description,
                        'threshold': threshold,
                        'is_active': True
                    })

                except Exception as row_error:
                    errors.append({
                        'row': row_idx,
                        'error': f"Error processing row: {str(row_error)}"
                    })
                    continue

            if not products and not errors:
                raise ValueError("No valid products found in Excel file")

            # Return data with processing results
            return {
                'success': True,
                'products': products,
                'errors': errors,
                'message': f"Processed {len(products) + len(errors)} rows, {len(products)} products ready for creation, {len(errors)} errors"
            }

        except Exception as e:
            self.logger.error(f"Excel processing error: {str(e)}")
            return {
                'success': False,
                'error': str(e),
                'message': 'Failed to process Excel file'
            }

    def _validate_headers(self, headers, required_headers):
        """Validate Excel has required columns"""
        missing = [h for h in required_headers if h not in headers]
        if missing:
            raise ValueError(f"Missing required columns: {', '.join(missing)}")

    def _get_or_create_product(self, sku, name, price, category):
        """
        Check product by SKU only, create if not exists
        Uses model defaults for description, threshold, and is_active
        """
        try:
            # Try to get existing product by SKU
            product = productsModel.objects.get(sku=sku)
            return product, False
        except productsModel.DoesNotExist:
            # Create new product with Excel data
            # Let model handle defaults for description, threshold, is_active
            product = productsModel.objects.create(
                sku=sku,
                name=name,
                price=price,
                category=category
                # description, threshold, is_active will use model defaults
            )
            return product, True
